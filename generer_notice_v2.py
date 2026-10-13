"""Génère la notice Word v2 du projet."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from crypto import ALPHABET, ALPHABET_SIZE, _ALPHABET_FR, _ALPHABET_V1

OUTPUT = Path(__file__).parent / "Notice_chiffrement_v2.docx"


def add_heading_styled(document, text, level=1):
    heading = document.add_heading(text, level=level)
    return heading


def add_para(document, text, bold=False):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)
    return paragraph


def main():
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Notice d'utilisation — Plateforme de chiffrement")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Version 2 — support du français (accents, cédille, ligatures)")
    run.italic = True
    run.font.size = Pt(12)

    add_heading_styled(document, "1. Objet", 1)
    add_para(
        document,
        "Cette notice décrit l'usage de la plateforme Streamlit et des scripts "
        "console associés. Le chiffrement est pédagogique : une substitution "
        "cyclique guidée par une clé (proche d'un chiffre de Vigenère sur un "
        "alphabet étendu). Il ne remplace pas un algorithme cryptographique "
        "moderne (AES, etc.).",
    )

    add_heading_styled(document, "2. Nouveautés de la version 2", 1)
    add_para(
        document,
        "La version 1 ignorait les lettres accentuées : é, à, ç, etc. étaient "
        "remplacés par un espace, ce qui rendait le texte français inutilisable. "
        "La version 2 conserve ces caractères.",
    )
    document.add_paragraph(
        "Lettres accentuées minuscules et majuscules : à â ä é è ê ë ï î ô ö ù û ü ÿ "
        "et À Â Ä É È Ê Ë Ï Î Ô Ö Ù Û Ü Ÿ.",
        style="List Bullet",
    )
    document.add_paragraph(
        "Cédille : ç et Ç.",
        style="List Bullet",
    )
    document.add_paragraph(
        "Ligatures françaises : œ/Œ (cœur, œuvre) et æ/Æ.",
        style="List Bullet",
    )
    document.add_paragraph(
        "Ponctuation typographique : « » ‘ ’ “ ” … — – ainsi que l'euro (€) et le degré (°).",
        style="List Bullet",
    )
    document.add_paragraph(
        "Sauts de ligne, tabulations et espaces insécables (copie depuis Word).",
        style="List Bullet",
    )
    add_para(
        document,
        f"L'alphabet compte désormais {ALPHABET_SIZE} caractères "
        f"({len(_ALPHABET_V1)} conservés de la v1 + {len(_ALPHABET_FR)} ajouts). "
        "Un caractère hors alphabet n'est plus transformé en espace : l'opération "
        "est refusée avec un message d'erreur.",
    )
    warning = document.add_paragraph()
    run = warning.add_run(
        "Incompatibilité : un message chiffré avec la v1 ne peut plus être "
        "déchiffré avec la v2. Rechiffrez les textes encore utiles."
    )
    run.bold = True

    add_heading_styled(document, "3. Alphabet pris en charge", 1)
    add_para(document, "Bloc hérité de la v1 (lettres ASCII, chiffres, ponctuation usuelle) :")
    document.add_paragraph(_ALPHABET_V1.replace("\n", "[LF]").replace("\t", "[TAB]"))
    add_para(document, "Bloc français et typographique ajouté en v2 :")
    visible_fr = (
        _ALPHABET_FR.replace("\n", " [saut de ligne] ")
        .replace("\t", " [tabulation] ")
        .replace("\u00a0", " [espace insécable] ")
    )
    document.add_paragraph(visible_fr)

    add_heading_styled(document, "4. Principe de l'algorithme", 1)
    add_para(
        document,
        "Chaque caractère est converti en un indice dans l'alphabet. Pour le "
        "chiffrement, on soustrait l'indice du caractère de clé correspondant "
        "(la clé est répétée cycliquement), modulo la taille de l'alphabet. "
        "Pour le déchiffrement, on additionne le même indice, toujours modulo "
        f"{ALPHABET_SIZE}. La clé doit faire de 1 à 13 caractères et appartenir "
        "au même alphabet.",
    )

    add_heading_styled(document, "5. Utiliser l'application Streamlit", 1)
    add_heading_styled(document, "5.1. Lancer en local", 2)
    document.add_paragraph("pip install -r requirements.txt", style="List Number")
    document.add_paragraph("streamlit run app.py", style="List Number")
    document.add_paragraph(
        "Ouvrir l'adresse indiquée, en général http://localhost:8501.",
        style="List Number",
    )

    add_heading_styled(document, "5.2. Chiffrer", 2)
    document.add_paragraph('Onglet « Chiffrer ».', style="List Number")
    document.add_paragraph(
        "Saisir une clé secrète (exemple : CléÉté56) puis le message en français.",
        style="List Number",
    )
    document.add_paragraph(
        "Cliquer sur « Chiffrer le message », puis copier le résultat.",
        style="List Number",
    )

    add_heading_styled(document, "5.3. Déchiffrer", 2)
    document.add_paragraph('Onglet « Déchiffrer ».', style="List Number")
    document.add_paragraph(
        "Saisir exactement la même clé, coller le texte chiffré, puis lancer le déchiffrement.",
        style="List Number",
    )

    add_heading_styled(document, "6. Scripts console", 1)
    add_para(
        document,
        "Chiffrement_intégral.py et Déchiffrement_intégral.py proposent le même "
        "algorithme v2 en ligne de commande. Ils importent le module crypto.py.",
    )

    add_heading_styled(document, "7. Déployer sur Streamlit Community Cloud", 1)
    document.add_paragraph(
        "Créer un compte GitHub si besoin, puis pousser ce projet sur un dépôt GitHub.",
        style="List Number",
    )
    document.add_paragraph(
        "Aller sur https://share.streamlit.io et se connecter avec GitHub.",
        style="List Number",
    )
    document.add_paragraph(
        "Cliquer sur Create app, puis Yup, I have an app.",
        style="List Number",
    )
    document.add_paragraph(
        "Sélectionner le dépôt, la branche main et le fichier d'entrée app.py.",
        style="List Number",
    )
    document.add_paragraph(
        "Laisser Python 3.12 si proposé, puis Deploy. L'URL sera du type https://….streamlit.app.",
        style="List Number",
    )
    document.add_paragraph(
        "Chaque nouveau git push met à jour l'application automatiquement.",
        style="List Number",
    )

    add_heading_styled(document, "8. Bonnes pratiques", 1)
    document.add_paragraph("Ne jamais partager la clé avec le message chiffré.", style="List Bullet")
    document.add_paragraph(
        "Ne pas utiliser cet outil pour des secrets professionnels ou médicaux.",
        style="List Bullet",
    )
    document.add_paragraph(
        "Tester un aller-retour (chiffrer puis déchiffrer) avant un usage réel.",
        style="List Bullet",
    )

    add_heading_styled(document, "9. Exemple", 1)
    add_para(document, "Message : « Déjà, le cœur français s'élance ! »")
    add_para(document, "Clé : Été2026")
    add_para(
        document,
        "Après chiffrement, le texte est illisible ; le déchiffrement avec la "
        "même clé doit restituer exactement le message, accents compris.",
    )

    document.save(OUTPUT)
    print(f"Notice écrite : {OUTPUT}")


if __name__ == "__main__":
    main()
